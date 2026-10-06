#!/usr/bin/env python3
"""Genera el documento AA4 (infografía comparativa) en DOCX con fuente NEGRA.

Pensado para editar y luego exportar a PDF. Fondo blanco, texto negro y las
imágenes antes/después en paralelo por cada aspecto.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = "/home/jenifrutica/SENA/NitroRythm/My project/Docs"
CAP = os.path.join(DOCS, "Capturas")
OUT = os.path.join(DOCS, "Evidencias", "GA5-AA4-EV01_Infografia_Comparativa.docx")

BLACK = RGBColor(0, 0, 0)


def black_run(paragraph, text, size=12, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.color.rgb = BLACK
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    return run


def heading(doc, text, level=1):
    h = doc.add_heading("", level=level)
    black_run(h, text, size=16 if level == 1 else 13, bold=True)
    return h


def body(doc, text):
    p = doc.add_paragraph()
    black_run(p, text, size=12)
    return p


def pair(doc, before_label, before_text, after_label, after_text):
    p = doc.add_paragraph()
    black_run(p, before_label, bold=True)
    black_run(p, before_text)
    p2 = doc.add_paragraph()
    black_run(p2, after_label, bold=True)
    black_run(p2, after_text)


def images(doc, before_img, after_img):
    table = doc.add_table(rows=1, cols=2)
    cells = table.rows[0].cells
    for cell, img in zip(cells, (before_img, after_img)):
        para = cell.paragraphs[0]
        if os.path.exists(img):
            run = para.add_run()
            run.add_picture(img, width=Inches(3.0))
        else:
            black_run(para, f"[falta {img}]", size=9)


def main():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    black_run(title, "NITRO RHYTHM", size=26, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    black_run(sub, "Infografía comparativa — antes / después de los ajustes", size=14, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    black_run(meta, "GA5-220501087-AA4-EV01 · SENA · Jenifer Daniela Urbano Córdoba", size=11)

    body(doc, "Comparación del prototipo NitroRhythm antes y después de los ajustes "
              "realizados a partir de la prueba de funcionalidad (5 usuarios). "
              "El estado «antes» se reconstruye a partir del prototipo inicial "
              "(primitivas y HUD en texto plano); el estado «después» son capturas "
              "reales de la versión final.")

    heading(doc, "1. Estructura e interfaces")
    pair(doc, "Antes: ", "una sola escena con menú y HUD en texto plano, sin navegación entre pantallas.",
         "Después: ", "cuatro escenas (menú, selección, gameplay y resultados) sobre Canvas + TextMeshPro, con HUD ampliado.")
    images(doc, os.path.join(CAP, "Antes/legacy/legacy_menu.png"), os.path.join(CAP, "Despues/ronda2/ui/ui_menu_principal.png"))

    heading(doc, "2. Concepto gráfico")
    pair(doc, "Antes: ", "karts, pista e interfaz con primitivas cúbicas de color plano; sin garaje 3D ni audio reactivo.",
         "Después: ", "karts y pilotos 3D texturizados, garaje 3D con pedestales y estética neón URP (bloom).")
    images(doc, os.path.join(CAP, "Antes/legacy/legacy_select.png"), os.path.join(CAP, "Despues/ronda3/seleccion_personajes.png"))

    heading(doc, "3. Niveles")
    pair(doc, "Antes: ", "pista de un solo estilo: plataformas grises de cubos, sin temática ni ambientación.",
         "Después: ", "siete dominios jugables, cada uno con cielo, props, partículas y música propios; checkpoints visibles.")
    images(doc, os.path.join(CAP, "Antes/legacy/legacy_track.png"), os.path.join(CAP, "Despues/ronda2/RESUMEN_dominios.png"))

    heading(doc, "4. Mecánicas")
    pair(doc, "Antes: ", "karts y villano de cubos; movimiento básico y sin combate.",
         "Después: ", "cuatro modos, villano que ataca al compás, K.O. con cuenta regresiva y feedback de impacto.")
    images(doc, os.path.join(CAP, "Antes/legacy/legacy_gameplay.png"), os.path.join(CAP, "Despues/ronda3/gameplay_ko_cuenta_regresiva.png"))

    heading(doc, "Conclusión")
    body(doc, "Los hallazgos P4 (HUD), P6 (feedback de impacto) y P9 (progresión) quedaron "
              "atendidos e implementados en la versión final y verificados con 40 pruebas "
              "automáticas (36 PlayMode + 4 EditMode).")

    doc.save(OUT)
    print("DOCX generado:", OUT)


if __name__ == "__main__":
    main()
